"""Transcript ingestion contracts and local file parsers."""

import hashlib
import html
import re
from datetime import datetime, timezone
from email import policy
from email.parser import BytesParser
from enum import StrEnum
from io import BytesIO
from pathlib import Path
from typing import Any

from pydantic import BaseModel, Field

MAX_TRANSCRIPT_BYTES = 10 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".txt", ".vtt", ".srt", ".docx", ".eml"}
TIMECODE = re.compile(
    r"^\s*(?:\d{2}:)?\d{2}:\d{2}[,.]\d{3}\s+-->\s+(?:\d{2}:)?\d{2}:\d{2}[,.]\d{3}"
)
VOICE_TAG = re.compile(r"^<v(?:\.[^ >]+)*(?:\s+([^>]+))?>(.*?)</v>$", re.IGNORECASE)
CAPTION_TAG = re.compile(r"<[^>]+>")


class TranscriptSource(StrEnum):
    UPLOAD = "Upload"
    MICROSOFT_365 = "Microsoft 365"
    TEAMS = "Teams"
    AZURE_BLOB = "Azure Blob"
    GOOGLE_DRIVE = "Google Drive"
    ZOOM = "Zoom"
    EMAIL = "Email"


class TranscriptDocument(BaseModel):
    source: TranscriptSource
    external_id: str
    title: str = Field(min_length=1)
    filename: str
    content_type: str | None = None
    text: str = Field(min_length=1)
    content_hash: str
    imported_at: datetime
    metadata: dict[str, Any] = Field(default_factory=dict)

    @classmethod
    def create(
        cls,
        *,
        source: TranscriptSource,
        external_id: str,
        title: str,
        filename: str,
        text: str,
        raw_content: bytes,
        content_type: str | None = None,
        metadata: dict[str, Any] | None = None,
    ) -> "TranscriptDocument":
        normalized = text.strip()
        if not normalized:
            raise ValueError("Transcript contains no readable text")
        return cls(
            source=source,
            external_id=external_id,
            title=title,
            filename=filename,
            content_type=content_type,
            text=normalized,
            content_hash=hashlib.sha256(raw_content).hexdigest(),
            imported_at=datetime.now(timezone.utc),
            metadata=metadata or {},
        )


def parse_uploaded_transcript(
    filename: str,
    content: bytes,
    content_type: str | None = None,
) -> TranscriptDocument:
    if len(content) > MAX_TRANSCRIPT_BYTES:
        raise ValueError("Transcript file exceeds the 10 MB limit")
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported transcript type. Supported types: {supported}")

    title = Path(filename).stem
    metadata: dict[str, Any] = {}
    if suffix in {".txt", ".vtt", ".srt"}:
        decoded = content.decode("utf-8-sig")
        text = _parse_caption_text(decoded) if suffix in {".vtt", ".srt"} else decoded
    elif suffix == ".docx":
        text = _parse_docx(content)
    else:
        title, text, metadata = _parse_eml(content, title)

    return TranscriptDocument.create(
        source=TranscriptSource.UPLOAD,
        external_id=hashlib.sha256(content).hexdigest(),
        title=title,
        filename=filename,
        text=text,
        raw_content=content,
        content_type=content_type,
        metadata=metadata,
    )


def _parse_caption_text(text: str) -> str:
    output: list[str] = []
    for block in re.split(r"\n\s*\n", text.replace("\r\n", "\n")):
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if not lines or lines[0] == "WEBVTT":
            continue
        if lines[0].startswith(("NOTE", "STYLE", "REGION")):
            continue

        timecode_index = next(
            (index for index, line in enumerate(lines) if TIMECODE.match(line)),
            None,
        )
        if timecode_index is None:
            continue
        payload = " ".join(lines[timecode_index + 1 :]).strip()
        voice = VOICE_TAG.match(payload)
        if voice:
            speaker, spoken_text = voice.groups()
            spoken_text = CAPTION_TAG.sub("", spoken_text).strip()
            speaker = speaker.strip() if speaker else None
            line = f"{speaker}: {spoken_text}" if speaker else spoken_text
        else:
            speaker = None
            line = CAPTION_TAG.sub("", payload).strip()
        line = html.unescape(line)
        if not line:
            continue
        speaker_prefix = f"{speaker}: " if speaker else None
        if speaker_prefix and output and output[-1].startswith(speaker_prefix):
            output[-1] = f"{output[-1]} {spoken_text}"
            continue
        if not output or output[-1] != line:
            output.append(line)
    return "\n".join(output)


def _parse_docx(content: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(content))
    return "\n".join(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())


def _parse_eml(content: bytes, fallback_title: str) -> tuple[str, str, dict[str, Any]]:
    message = BytesParser(policy=policy.default).parsebytes(content)
    title = str(message.get("subject") or fallback_title)
    metadata = {
        "from": str(message.get("from") or ""),
        "date": str(message.get("date") or ""),
    }
    for attachment in message.iter_attachments():
        attachment_name = attachment.get_filename() or ""
        if Path(attachment_name).suffix.lower() in SUPPORTED_EXTENSIONS - {".eml"}:
            payload = attachment.get_payload(decode=True) or b""
            nested = parse_uploaded_transcript(attachment_name, payload, attachment.get_content_type())
            metadata["attachment"] = attachment_name
            return title, nested.text, metadata

    body = message.get_body(preferencelist=("plain",))
    if body:
        return title, body.get_content(), metadata
    raise ValueError("Email contains no supported transcript attachment or plain-text body")