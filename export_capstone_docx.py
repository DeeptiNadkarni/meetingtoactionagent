"""Export the capstone Markdown document as a styled Microsoft Word file."""

import argparse
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DEFAULT_SOURCE = Path("CAPSTONE_DOCUMENTATION.md")
DEFAULT_OUTPUT = Path("Meeting-to-Action-Capstone-Documentation-Updated.docx")


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_inline_text(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def add_paragraph(document: Document, text: str, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    add_inline_text(paragraph, text)


def configure_document(document: Document, document_label: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    heading_sizes = {"Title": 28, "Heading 1": 20, "Heading 2": 15, "Heading 3": 12}
    for style_name, size in heading_sizes.items():
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(8, 127, 115)
        style.font.bold = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"Meeting to Action | {document_label}")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(80, 99, 95)


def add_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        add_inline_text(cell.paragraphs[0], value)
        set_cell_shading(cell, "DCEFEB")
        for run in cell.paragraphs[0].runs:
            run.bold = True

    for source_row in rows[2:]:
        cells = table.add_row().cells
        for index, value in enumerate(source_row):
            add_inline_text(cells[index].paragraphs[0], value)
    document.add_paragraph()


def export(
    source: Path = DEFAULT_SOURCE,
    output: Path = DEFAULT_OUTPUT,
    template: Path | None = None,
) -> None:
    document_label = (
        "System Design" if source.name.casefold() == "system_design.md" else "Capstone Project Documentation"
    )
    document = Document(template) if template else Document()
    configure_document(document, document_label)
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_lines:
            add_paragraph(document, " ".join(paragraph_lines))
            paragraph_lines.clear()

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if stripped == '<div style="page-break-after: always;"></div>':
            flush_paragraph()
            document.add_page_break()
            index += 1
            continue

        if stripped.startswith("```"):
            flush_paragraph()
            language = stripped[3:].strip()
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            run = paragraph.add_run("Architecture flow\n" if language == "mermaid" else "")
            run.bold = True
            code_run = paragraph.add_run("\n".join(code_lines))
            code_run.font.name = "Consolas"
            code_run.font.size = Pt(8.5)
            index += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            flush_paragraph()
            table_lines: list[list[str]] = []
            while index < len(lines):
                candidate = lines[index].strip()
                if not (candidate.startswith("|") and candidate.endswith("|")):
                    break
                table_lines.append([cell.strip() for cell in candidate.strip("|").split("|")])
                index += 1
            add_table(document, table_lines)
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline_text(paragraph, text)
            else:
                add_paragraph(document, text, f"Heading {level - 1}")
            index += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        bullet = re.match(r"^-\s+(.+)$", stripped)
        if numbered or bullet:
            flush_paragraph()
            add_paragraph(
                document,
                (numbered or bullet).group(2 if numbered else 1),
                "List Number" if numbered else "List Bullet",
            )
            index += 1
            continue

        if not stripped or stripped == "---":
            flush_paragraph()
            index += 1
            continue

        paragraph_lines.append(stripped)
        index += 1

    flush_paragraph()
    document.core_properties.title = f"Meeting to Action - {document_label}"
    document.core_properties.subject = "AI capstone architecture, evaluation, and testing"
    document.core_properties.author = "Meeting to Action Project"
    document.save(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--template", type=Path)
    arguments = parser.parse_args()
    export(arguments.source, arguments.output, arguments.template)